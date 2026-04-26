"""
Generates OnChainCarFacts-Whitepaper.docx from docs/whitepaper.md.

Minimal Markdown -> docx renderer sufficient for this paper's subset:
headings, paragraphs, bullet/numbered lists, tables, inline **bold**,
*italic*, and `code`. Not a general-purpose converter.
"""

import os
import re

from docx import Document
from docx.shared import Pt, RGBColor, Inches

HERE = os.path.dirname(os.path.abspath(__file__))
INPUT = os.path.join(HERE, "whitepaper.md")
OUTPUT = os.path.join(HERE, "OnChainCarFacts-Whitepaper.docx")

MONO = "Consolas"

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_inline(paragraph, text):
    """Render inline Markdown (bold/italic/code) into paragraph runs."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = MONO
            r.font.size = Pt(10)
        elif part.startswith("*") and part.endswith("*"):
            r = paragraph.add_run(part[1:-1])
            r.italic = True
        else:
            paragraph.add_run(part)


def parse_table_row(line):
    # strip leading/trailing | and split
    inner = line.strip()
    if inner.startswith("|"):
        inner = inner[1:]
    if inner.endswith("|"):
        inner = inner[:-1]
    return [c.strip() for c in inner.split("|")]


def build():
    with open(INPUT, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.rstrip()

        # Blank line -> just advance
        if not stripped.strip():
            i += 1
            continue

        # Horizontal rule -> skip (we use headings for structure)
        if stripped.strip() == "---":
            i += 1
            continue

        # Headings
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1)) - 1  # # -> 0 (title), ## -> 1, ### -> 2
            text = m.group(2).strip()
            h = doc.add_heading(level=min(level, 4))
            add_inline(h, text)
            i += 1
            continue

        # Table: current line starts with "|" and next line is the separator
        if stripped.startswith("|") and i + 1 < n and re.match(
            r"^\|[\s\-:|]+\|\s*$", lines[i + 1]
        ):
            headers = parse_table_row(stripped)
            i += 2  # skip header + separator
            rows = []
            while i < n and lines[i].lstrip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            t = doc.add_table(rows=1 + len(rows), cols=len(headers))
            t.style = "Light Grid Accent 1"
            hdr_cells = t.rows[0].cells
            for ci, h in enumerate(headers):
                hdr_cells[ci].text = ""
                p = hdr_cells[ci].paragraphs[0]
                r = p.add_run(h)
                r.bold = True
            for ri, row in enumerate(rows, start=1):
                cells = t.rows[ri].cells
                for ci, val in enumerate(row):
                    cells[ci].text = ""
                    add_inline(cells[ci].paragraphs[0], val)
            continue

        # Numbered list
        m_num = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m_num:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m_num.group(2))
            i += 1
            # consume continuation lines indented with spaces
            while i < n and re.match(r"^   +\S", lines[i]):
                cont = lines[i].strip()
                p.add_run(" ")
                add_inline(p, cont)
                i += 1
            continue

        # Bulleted list
        m_bul = re.match(r"^[-*]\s+(.*)$", stripped)
        if m_bul:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m_bul.group(1))
            i += 1
            while i < n and re.match(r"^   +\S", lines[i]):
                cont = lines[i].strip()
                p.add_run(" ")
                add_inline(p, cont)
                i += 1
            continue

        # Blockquote
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            r_run = p.add_run()
            add_inline(p, stripped[2:])
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # Paragraph: accumulate consecutive non-blank, non-special lines
        buf = [stripped]
        i += 1
        while i < n:
            nxt = lines[i]
            if not nxt.strip():
                break
            if re.match(r"^#{1,4}\s", nxt):
                break
            if nxt.lstrip().startswith("|"):
                break
            if re.match(r"^\d+\.\s", nxt) or re.match(r"^[-*]\s", nxt):
                break
            if nxt.strip() == "---":
                break
            buf.append(nxt.strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(buf))

    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    build()
